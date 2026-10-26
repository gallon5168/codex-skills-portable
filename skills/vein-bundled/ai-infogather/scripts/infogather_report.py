#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
infogather_report.py — 被动信息收集 xlsx 报告生成器

用法:
    # 直接输出空模板到桌面
    python.exe infogather_report.py --target "目标名称" --output 桌面路径/report.xlsx

    # 从 JSON 文件导入已有数据
    python.exe infogather_report.py --target "目标名称" --input data.json --output report.xlsx

JSON 输入格式:
    {
        "domains": [{"domain":"...", "unit":"...", "ip":"...", "source":"...", "note":"..."}],
        "ips": [{"ip":"...", "domain":"...", "location":"...", "isp":"...", "type":"...", "source":"..."}],
        "urls": [{"url":"...", "system":"...", "accessible":"...", "ip":"...", "source":"..."}],
        "systems": [{"system":"...", "tech":"...", "unit":"...", "ports":"...", "source":"..."}],
        "social": [{"type":"...", "content":"...", "belong":"...", "source":"..."}],
        "todos": [{"priority":"...", "action":"..."}]
    }
"""

import json
import sys
import os
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[!] 请先安装 openpyxl: pip install openpyxl")
    sys.exit(1)

# ── 样式定义 ──
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_FONT = Font(name="微软雅黑", size=10)
BODY_ALIGN = Alignment(vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="B4C6E7"),
    right=Side(style="thin", color="B4C6E7"),
    top=Side(style="thin", color="B4C6E7"),
    bottom=Side(style="thin", color="B4C6E7"),
)

ALT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

# 优先级颜色
PRIO_FILLS = {
    "⭐⭐⭐": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
    "⭐⭐": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
    "⭐": PatternFill(start_color="6BCB77", end_color="6BCB77", fill_type="solid"),
}


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def style_body(ws, start_row, end_row, cols):
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            cell.border = THIN_BORDER
            if (r - start_row) % 2 == 1:
                cell.fill = ALT_FILL


def auto_width(ws, cols, min_width=10, max_width=40):
    for c in range(1, cols + 1):
        letter = get_column_letter(c)
        lengths = []
        for row in ws.iter_rows(min_col=c, max_col=c, values_only=True):
            val = row[0]
            if val:
                # 中文按2倍宽度估算
                clen = sum(2 if ord(ch) > 127 else 1 for ch in str(val))
                lengths.append(clen)
        best = max(lengths) if lengths else min_width
        ws.column_dimensions[letter].width = min(max(best + 2, min_width), max_width)


# ── Sheet 定义 ──
SHEETS = [
    {
        "name": "股权结构",
        "headers": ["层级", "公司名称", "持股比例", "注册资本", "法人", "成立日期", "参保数", "官网/域名", "ICP备案", "邮箱", "地址", "来源"],
        "widths": [10, 35, 12, 18, 12, 14, 10, 25, 20, 35, 35, 15],
        "key": "equity",
        "fields": ["level", "company", "share", "capital", "person", "date", "insured", "website", "icp", "emails", "address", "source"],
    },
    {
        "name": "域名资产",
        "headers": ["域名", "所属单位", "IP", "来源", "备注"],
        "widths": [30, 30, 18, 20, 30],
        "key": "domains",
        "fields": ["domain", "unit", "ip", "source", "note"],
    },
    {
        "name": "IP资产",
        "headers": ["IP", "关联域名", "地理位置", "运营商", "类型", "来源"],
        "widths": [18, 30, 18, 12, 12, 20],
        "key": "ips",
        "fields": ["ip", "domain", "location", "isp", "type", "source"],
    },
    {
        "name": "URL资产",
        "headers": ["URL", "系统类型", "可访问性", "IP", "来源", "备注"],
        "widths": [50, 20, 12, 18, 20, 30],
        "key": "urls",
        "fields": ["url", "system", "accessible", "ip", "source", "note"],
    },
    {
        "name": "系统资产",
        "headers": ["系统", "技术栈", "归属单位", "端口猜测", "来源"],
        "widths": [30, 25, 25, 15, 25],
        "key": "systems",
        "fields": ["system", "tech", "unit", "ports", "source"],
    },
    {
        "name": "社工线索",
        "headers": ["类型", "内容", "归属", "来源"],
        "widths": [15, 40, 25, 25],
        "key": "social",
        "fields": ["type", "content", "belong", "source"],
    },
    {
        "name": "待办列表",
        "headers": ["优先级", "操作建议"],
        "widths": [12, 80],
        "key": "todos",
        "fields": ["priority", "action"],
    },
]


def build_sheet(ws, data, sheet_def):
    """填充一个 sheet 的数据"""
    headers = sheet_def["headers"]
    fields = sheet_def["fields"]
    rows = data.get(sheet_def["key"], [])

    # 写表头
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))

    # 写数据
    for r, row_data in enumerate(rows, 2):
        for c, field in enumerate(fields, 1):
            val = row_data.get(field, "")
            cell = ws.cell(row=r, column=c, value=val)
            # 优先级着色
            if field == "priority" and val in PRIO_FILLS:
                cell.fill = PRIO_FILLS[val]
                cell.font = Font(name="微软雅黑", bold=True, size=10, color="FFFFFF")

    if rows:
        style_body(ws, 2, len(rows) + 1, len(headers))
        auto_width(ws, len(headers))


def generate_report(target, output_path, data=None):
    """生成 xlsx 报告"""
    if data is None:
        data = {}

    wb = Workbook()
    # 删除默认 sheet
    wb.remove(wb.active)

    for sheet_def in SHEETS:
        ws = wb.create_sheet(title=sheet_def["name"])
        build_sheet(ws, data, sheet_def)

    # 在第一个 sheet 前插入封面信息
    ws = wb.create_sheet(title="概览", index=0)
    cover = [
        ("infogather 被动信息收集报告", ""),
        ("", ""),
        ("目标名称", target),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("工具", "infogather skill + Claude Code"),
        ("", ""),
    ]
    for name in SHEETS:
        n = len(data.get(name["key"], []))
        cover.append((f"  {name['name']}", f"{n} 条"))

    for r, (k, v) in enumerate(cover, 1):
        cell_k = ws.cell(row=r, column=1, value=k)
        cell_v = ws.cell(row=r, column=2, value=v)
        if r == 1:
            cell_k.font = Font(name="微软雅黑", bold=True, size=16)
            cell_v.font = Font(name="微软雅黑", size=14)
        elif k:
            cell_k.font = Font(name="微软雅黑", bold=True, size=11)
            cell_v.font = Font(name="微软雅黑", size=11)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40

    wb.save(output_path)
    return output_path


def create_empty_template(target, output_path):
    """创建空模板，供手动填写"""
    data = {
        "domains": [],
        "ips": [],
        "urls": [],
        "systems": [],
        "social": [],
        "todos": [
            {"priority": "⭐⭐⭐", "action": f"nslookup 确认 {target} 的主域名解析 IP"},
            {"priority": "⭐⭐⭐", "action": f"crt.sh 查询证书子域名"},
            {"priority": "⭐⭐⭐", "action": "子域名爆破（subfinder / oneforall）"},
            {"priority": "⭐⭐", "action": "FOFA/Hunter 搜 domain/icp.name"},
            {"priority": "⭐⭐", "action": "已知 IP 全端口扫描"},
            {"priority": "⭐⭐", "action": "招标公告挖掘行业系统"},
            {"priority": "⭐", "action": "GitHub/语雀/网盘泄露搜索"},
        ],
    }
    return generate_report(target, output_path, data)


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="infogather xlsx 报告生成器")
    parser.add_argument("--target", required=True, help="目标名称")
    parser.add_argument("--output", default=None, help="输出 xlsx 路径")
    parser.add_argument("--input", default=None, help="输入 JSON 数据文件")
    args = parser.parse_args()

    # 默认输出到桌面
    if not args.output:
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
        safe_name = args.target.replace(" ", "_").replace("（", "").replace("）", "")[:20]
        args.output = os.path.join(desktop, f"infogather_{safe_name}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")

    data = None
    if args.input:
        data = load_json(args.input)

    if data:
        path = generate_report(args.target, args.output, data)
    else:
        path = create_empty_template(args.target, args.output)

    print(f"[+] 报告已生成: {path}")
    print(f"[+] 共 {len(data or {})} 个数据源" if data else "[+] 空模板已创建，请手动填写数据")
